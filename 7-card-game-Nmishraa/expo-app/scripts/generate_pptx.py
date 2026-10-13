import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Define design color tokens
DARK_BG = RGBColor(7, 22, 12)       # Rich dark emerald dark mode background (#07160C)
NAVY_BG = RGBColor(10, 22, 40)     # Dark slate navy container panel (#0A1628)
CARD_BG = RGBColor(15, 29, 48)     # Card panel fill (#0F1D30)
EMERALD_GREEN = RGBColor(22, 163, 74) # Accent green (#16A34A)
GOLD_ACCENT = RGBColor(234, 179, 8)   # Accent gold (#EAB308)
CYAN_ACCENT = RGBColor(56, 189, 248)  # Accent cyan (#38BDF8)
TEXT_WHITE = RGBColor(255, 255, 255)  # Primary text (#FFFFFF)
TEXT_LIGHT = RGBColor(203, 213, 225)  # Secondary text (#CBD5E1)
TEXT_MUTED = RGBColor(148, 163, 184)  # Muted text (#94A3B8)
BORDER_COLOR = RGBColor(30, 58, 90)   # Subtle panel border

ASSETS_DIR = os.path.join(os.path.dirname(__file__), 'assets_presentation')
OUTPUT_PPTX = os.path.join(os.path.dirname(__file__), '..', '7-Cards-Least-Game-Presentation.pptx')

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank layout

    def add_slide_background(slide, color=DARK_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.color.rgb = color
        return bg

    def add_header(slide, title_text, category_text="7 CARDS LEAST"):
        # Header Badge / Category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = GOLD_ACCENT

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    def add_card_shape(slide, left, top, width, height, bg_color=NAVY_BG, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        return shape

    # =========================================================================
    # SLIDE 1: Hero / Cover
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide1, DARK_BG)

    # Decorative background card graphic box
    add_card_shape(slide1, Inches(0.8), Inches(1.2), Inches(11.733), Inches(5.3), NAVY_BG, EMERALD_GREEN)

    # Title & Subtitle Left Block
    tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(6.0), Inches(4.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "🎴 7 CARDS LEAST"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "Play Smart.\nKeep Your Score Low.\nHave More Fun."
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.space_after = Pt(18)

    p3 = tf1.add_paragraph()
    p3.text = "The fast-paced online shedding card game where lowest score wins. Play solo against computer AI or host private multiplayer rooms with friends."
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_LIGHT
    p3.space_after = Pt(24)

    p4 = tf1.add_paragraph()
    p4.text = "🌐 PLAY FREE ONLINE: cards.gnanamai.com"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = CYAN_ACCENT

    # Screenshot Right Mockup Frame
    img1_path = os.path.join(ASSETS_DIR, 'home_screen.png')
    if os.path.exists(img1_path):
        slide1.shapes.add_picture(img1_path, Inches(7.3), Inches(1.8), width=Inches(4.8))

    slide1.notes_slide.notes_text_frame.text = (
        "Slide 1 Talking Points:\n"
        "- Welcome to the 7 Cards Least presentation.\n"
        "- 7 Cards Least is a modern, fast-paced online shedding card game hosted on cards.gnanamai.com.\n"
        "- Players aim to shed high-value cards, form matching combinations, and finish with the lowest score.\n"
        "- Available 100% free online directly in any browser—no app downloads required."
    )

    # =========================================================================
    # SLIDE 2: Meet the Game
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide2, DARK_BG)
    add_header(slide2, "Meet 7 Cards Least: The Lowest Score Card Game")

    # Left 3 Pillar Cards
    pillars = [
        ("🎯 Objective: Lowest Score", "Unlike traditional card games where accumulating high points wins, 7 Cards Least flips the mechanics: finish with the LEAST points to win."),
        ("⚡ 100% Free & Browser-Based", "Play instantly on desktop, mobile, or tablet. No app store downloads, logins, or installation friction required."),
        ("🤖 Solo AI & Multiplayer", "Practice anytime against 3 smart computer bots (AlphaBot, BetaBot, OmegaBot) or create private 4-digit code tables for friends.")
    ]

    for i, (p_title, p_desc) in enumerate(pillars):
        top_pos = Inches(1.6 + i * 1.7)
        add_card_shape(slide2, Inches(0.8), top_pos, Inches(5.8), Inches(1.5), NAVY_BG, BORDER_COLOR)
        tb = slide2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.15), Inches(5.4), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = p_title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = GOLD_ACCENT
        p1.space_after = Pt(6)
        p2 = tf.add_paragraph()
        p2.text = p_desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_LIGHT

    # Right Screenshot
    if os.path.exists(img1_path):
        slide2.shapes.add_picture(img1_path, Inches(7.0), Inches(1.6), width=Inches(5.5))

    slide2.notes_slide.notes_text_frame.text = (
        "Slide 2 Talking Points:\n"
        "- Introduce 7 Cards Least as a classic shedding card game inverted for maximum fun.\n"
        "- Key objective: shed high face cards, preserve zero-point Jokers, and declare LEAST when hand points <= 10.\n"
        "- Highly accessible on web browsers across mobile and desktop."
    )

    # =========================================================================
    # SLIDE 3: How to Play
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide3, DARK_BG)
    add_header(slide3, "How to Play: Turn-by-Turn Gameplay Flow")

    steps = [
        ("1. Deal 7 Cards & Joker", "7 cards dealt per player. 1 face-up card sets the zero-point Joker rank for the round; 1 starts the Discard Pile."),
        ("2. Discard First", "Select 1 single card, equal rank set (e.g. 8-8-8), or suited run (e.g. 4♣-5♣-6♣) to drop on the Discard Pile."),
        ("3. Draw Next", "Draw 1 replacement card from either the face-down Draw Deck or the top face-up card of the Discard Pile."),
        ("4. Declare LEAST!", "When hand sum is ≤ 10 points, call LEAST! Lowest score wins round (0 pts). Equal or higher call incurs 80-pt penalty!")
    ]

    for i, (s_title, s_desc) in enumerate(steps):
        left_pos = Inches(0.8 + (i % 2) * 5.9)
        top_pos = Inches(1.6 + (i // 2) * 2.6)
        add_card_shape(slide3, left_pos, top_pos, Inches(5.6), Inches(2.3), NAVY_BG, EMERALD_GREEN if i == 3 else BORDER_COLOR)
        
        tb = slide3.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(0.2), Inches(5.2), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = s_title
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = GOLD_ACCENT if i == 3 else CYAN_ACCENT
        p1.space_after = Pt(8)
        p2 = tf.add_paragraph()
        p2.text = s_desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_LIGHT

    slide3.notes_slide.notes_text_frame.text = (
        "Slide 3 Talking Points:\n"
        "- Explain the exact turn flow: Discard first, then Draw next.\n"
        "- Highlighting the Match & Skip rule: dropping a card matching the discard top rank skips drawing.\n"
        "- Emphasize the Least call threshold (10 pts or less) and the severe 80-point wrong call penalty."
    )

    # =========================================================================
    # SLIDE 4: The Strategy
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide4, DARK_BG)
    add_header(slide4, "Strategic Card Shedding & Scoring Rules")

    # Left Column: Card Values
    add_card_shape(slide4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), NAVY_BG, BORDER_COLOR)
    tb_vals = slide4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf_vals = tb_vals.text_frame
    tf_vals.word_wrap = True
    
    p = tf_vals.paragraphs[0]
    p.text = "🎴 Card Point Values"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(12)

    val_points = [
        ("Ace (A):", "1 Point (Best standard low card)"),
        ("Cards 2 to 10:", "Face Value (2–10 penalty points)"),
        ("Face Cards (J, Q, K):", "10 Points each (Heavy point liability!)"),
        ("Table Joker Rank:", "0 Points (Wildcard setup rank)"),
        ("Least Call Threshold:", "Hand sum must be ≤ 10 Points"),
        ("False Call Penalty:", "80 Points added if opponent ≤ caller")
    ]
    for lbl, desc in val_points:
        p = tf_vals.add_paragraph()
        p.text = f"• {lbl} {desc}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

    # Right Column: Key Tactics
    add_card_shape(slide4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), NAVY_BG, EMERALD_GREEN)
    tb_tac = slide4.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf_tac = tb_tac.text_frame
    tf_tac.word_wrap = True

    p = tf_tac.paragraphs[0]
    p.text = "💡 Tactical Principles for Winning"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(12)

    tactics = [
        ("Early Face Card Dump:", "Discard J, Q, K in turns 1–3 to avoid getting caught with 10-point penalties when opponents call Least."),
        ("Multi-Card Combinations:", "Build matching rank sets or suited runs (e.g. 5♠-6♠-7♠) to shed 3–4 cards in a single move."),
        ("Zero-Point Joker Hold:", "Hold Table Jokers to guarantee low hand totals and execute unchallengeable Least declarations."),
        ("Discard Pile Tracking:", "Track what opponents pick from the open pile to estimate their hand points before declaring Least.")
    ]
    for lbl, desc in tactics:
        p = tf_tac.add_paragraph()
        p.text = f"• {lbl} {desc}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    slide4.notes_slide.notes_text_frame.text = (
        "Slide 4 Talking Points:\n"
        "- Detail the scoring breakdown: Aces = 1pt, Face cards = 10pts, Jokers = 0pts.\n"
        "- Explain strategic dumping of 10-point face cards early in the game.\n"
        "- Emphasize set building and opponent tracking to minimize hand score safely."
    )

    # =========================================================================
    # SLIDE 5: Play Against AI
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide5, DARK_BG)
    add_header(slide5, "Singleplayer Mode: Play Against Computer AI")

    # Left Info Column
    add_card_shape(slide5, Inches(0.8), Inches(1.6), Inches(4.8), Inches(5.2), NAVY_BG, BORDER_COLOR)
    tb_ai = slide5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(4.4), Inches(4.8))
    tf_ai = tb_ai.text_frame
    tf_ai.word_wrap = True

    p = tf_ai.paragraphs[0]
    p.text = "🤖 Smart AI Opponents"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(12)

    ai_features = [
        ("Intelligent Bot Logic:", "AlphaBot 🤖, BetaBot 🤖, and OmegaBot 🤖 evaluate optimal set combinations and card shedding."),
        ("Instant Solo Practice:", "Zero queue times or player waiting. Jump into a fast-paced game instantly in your browser."),
        ("Realistic Strategy:", "Bots track table discards, hold zero-point Jokers, and declare Least strategically."),
        ("Custom Round Limits:", "Play quick 1-round matches or extended 20-round tournaments.")
    ]
    for lbl, desc in ai_features:
        p = tf_ai.add_paragraph()
        p.text = f"• {lbl} {desc}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    # Right Screenshot
    img_ai = os.path.join(ASSETS_DIR, 'ai_gameplay.png')
    if os.path.exists(img_ai):
        slide5.shapes.add_picture(img_ai, Inches(5.9), Inches(1.6), width=Inches(6.6))

    slide5.notes_slide.notes_text_frame.text = (
        "Slide 5 Talking Points:\n"
        "- Highlight Singleplayer Mode for solo practice.\n"
        "- AI opponents simulate realistic competitive decision-making.\n"
        "- Instant match start with no waiting for human lobby fill."
    )

    # =========================================================================
    # SLIDE 6: Play with Friends & Family
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide6, DARK_BG)
    add_header(slide6, "Multiplayer: Play With Friends & Family")

    # Left Screenshot
    img_mp = os.path.join(ASSETS_DIR, 'multiplayer_screen.png')
    if os.path.exists(img_mp):
        slide6.shapes.add_picture(img_mp, Inches(0.8), Inches(1.6), width=Inches(6.5))

    # Right Info Column
    add_card_shape(slide6, Inches(7.6), Inches(1.6), Inches(4.9), Inches(5.2), NAVY_BG, EMERALD_GREEN)
    tb_mp = slide6.shapes.add_textbox(Inches(7.8), Inches(1.8), Inches(4.5), Inches(4.8))
    tf_mp = tb_mp.text_frame
    tf_mp.word_wrap = True

    p = tf_mp.paragraphs[0]
    p.text = "🎮 Real-Time Multiplayer"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(12)

    mp_features = [
        ("Private Room Codes:", "Create a private lobby and generate a 4-digit room code for instant access."),
        ("WhatsApp & Link Invites:", "Share direct room links via WhatsApp to invite friends and family in 1 click."),
        ("Custom Room Settings:", "Set round limits (1–20 rounds) and turn time limits (60s or unlimited)."),
        ("Live Table Chat:", "Send real-time chat messages and reactions during table play.")
    ]
    for lbl, desc in mp_features:
        p = tf_mp.add_paragraph()
        p.text = f"• {lbl} {desc}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    slide6.notes_slide.notes_text_frame.text = (
        "Slide 6 Talking Points:\n"
        "- Detail the online multiplayer private table creation.\n"
        "- 4-digit room codes and WhatsApp sharing for instant friend invites.\n"
        "- Support for 2 to 8 players with real-time socket state synchronization."
    )

    # =========================================================================
    # SLIDE 7: The Game Experience
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide7, DARK_BG)
    add_header(slide7, "The Game Experience: Interface & Table Mechanics")

    # Left Info Box
    add_card_shape(slide7, Inches(0.8), Inches(1.6), Inches(5.0), Inches(5.2), NAVY_BG, BORDER_COLOR)
    tb_exp = slide7.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(4.6), Inches(4.8))
    tf_exp = tb_exp.text_frame
    tf_exp.word_wrap = True

    p = tf_exp.paragraphs[0]
    p.text = "🌟 Premium Table UI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(12)

    exp_features = [
        ("Felt Table Themes:", "Select custom felt table colors (Classic Emerald Green, Midnight Dark, Velvet Red)."),
        ("Turn Banners & Animations:", "Clear phase guidance (⚡ YOUR TURN: DISCARD! / 🎴 YOUR TURN: PICK!) with smooth draw/discard card flight physics."),
        ("Hand Sorting Tool:", "One-touch 'Sort Cards' button automatically arranges hand by sequence rank and suit order."),
        ("Score Leaderboard:", "Real-time round score calculations and cumulative match total tracking.")
    ]
    for lbl, desc in exp_features:
        p = tf_exp.add_paragraph()
        p.text = f"• {lbl} {desc}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    # Right Screenshot (Lobby / Gameplay)
    img_lobby = os.path.join(ASSETS_DIR, 'lobby_screen.png')
    if os.path.exists(img_lobby):
        slide7.shapes.add_picture(img_lobby, Inches(6.1), Inches(1.6), width=Inches(6.4))

    slide7.notes_slide.notes_text_frame.text = (
        "Slide 7 Talking Points:\n"
        "- Showcase the polished game table UI, felt theme selection, and animations.\n"
        "- Highlight turn phase prompts and automatic hand sorting features."
    )

    # =========================================================================
    # SLIDE 8: Why Play 7 Cards Least?
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide8, DARK_BG)
    add_header(slide8, "Why Play 7 Cards Least? Key Value Pillars")

    pillars_why = [
        ("🌐 100% Free Web App", "Play instantly in any desktop, tablet, or mobile web browser. Zero app downloads, zero installation friction, zero hidden fees."),
        ("⚡ Thrilling Inverted Mechanics", "Reverses traditional card scoring: shedding high cards and managing hand points creates intense, strategic end-game decisions."),
        ("👥 Flexible Solo & Social Modes", "Switch seamlessly between fast singleplayer AI practice and private multiplayer rooms with family and friends."),
        ("🎨 Modern Dark Mode UI", "Polished dark emerald aesthetic, clear card typography, realistic card physics, and smooth sound effects.")
    ]

    for i, (w_title, w_desc) in enumerate(pillars_why):
        left_pos = Inches(0.8 + (i % 2) * 5.9)
        top_pos = Inches(1.6 + (i // 2) * 2.6)
        add_card_shape(slide8, left_pos, top_pos, Inches(5.6), Inches(2.3), NAVY_BG, EMERALD_GREEN if i == 0 else BORDER_COLOR)
        
        tb = slide8.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(0.2), Inches(5.2), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = w_title
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = GOLD_ACCENT
        p1.space_after = Pt(8)
        p2 = tf.add_paragraph()
        p2.text = w_desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_LIGHT

    slide8.notes_slide.notes_text_frame.text = (
        "Slide 8 Talking Points:\n"
        "- Summarize core value pillars: instant web accessibility, inverted shedding gameplay, and modern UI design."
    )

    # =========================================================================
    # SLIDE 9: Start Playing
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide9, DARK_BG)

    add_card_shape(slide9, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5), NAVY_BG, EMERALD_GREEN)

    tb9 = slide9.shapes.add_textbox(Inches(1.5), Inches(1.4), Inches(10.333), Inches(4.7))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    p = tf9.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "⚡ YOUR NEXT GAME STARTS HERE"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(12)

    p2 = tf9.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "cards.gnanamai.com"
    p2.font.size = Pt(38)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.space_after = Pt(24)

    # 2 Big Action Boxes side by side
    add_card_shape(slide9, Inches(1.8), Inches(3.4), Inches(4.5), Inches(2.3), DARK_BG, CYAN_ACCENT)
    tb_act1 = slide9.shapes.add_textbox(Inches(2.0), Inches(3.6), Inches(4.1), Inches(1.9))
    tf_act1 = tb_act1.text_frame
    tf_act1.word_wrap = True
    p = tf_act1.paragraphs[0]
    p.text = "🤖 Practice vs Computer AI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p.space_after = Pt(6)
    p = tf_act1.add_paragraph()
    p.text = "Start instant solo match against computer bots with zero wait time."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT

    add_card_shape(slide9, Inches(7.0), Inches(3.4), Inches(4.5), Inches(2.3), DARK_BG, EMERALD_GREEN)
    tb_act2 = slide9.shapes.add_textbox(Inches(7.2), Inches(3.6), Inches(4.1), Inches(1.9))
    tf_act2 = tb_act2.text_frame
    tf_act2.word_wrap = True
    p = tf_act2.paragraphs[0]
    p.text = "🎮 Play With Friends"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    p.space_after = Pt(6)
    p = tf_act2.add_paragraph()
    p.text = "Host a private table, generate a 4-digit room code, and share via WhatsApp."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT

    slide9.notes_slide.notes_text_frame.text = (
        "Slide 9 Talking Points:\n"
        "- Call to action directing audience to cards.gnanamai.com.\n"
        "- Highlight two immediate starting options: Solo AI practice or Private Multiplayer Host."
    )

    # =========================================================================
    # SLIDE 10: Closing
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide10, DARK_BG)

    add_card_shape(slide10, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1), NAVY_BG, GOLD_ACCENT)

    tb10 = slide10.shapes.add_textbox(Inches(2.0), Inches(1.8), Inches(9.333), Inches(3.9))
    tf10 = tb10.text_frame
    tf10.word_wrap = True

    p = tf10.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "🎴 7 CARDS LEAST"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(18)

    p2 = tf10.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "“Gather Your Players.\nPlay Your Cards.\nGo for the Lowest Score.”"
    p2.font.size = Pt(34)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.space_after = Pt(24)

    p3 = tf10.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.text = "Thank You! Play online at https://cards.gnanamai.com/"
    p3.font.size = Pt(16)
    p3.font.bold = True
    p3.font.color.rgb = CYAN_ACCENT

    slide10.notes_slide.notes_text_frame.text = (
        "Slide 10 Talking Points:\n"
        "- Final closing slide reinforcing the game slogan and brand web address."
    )

    prs.save(OUTPUT_PPTX)
    console_log(f"[PPTX Generator] Successfully generated PowerPoint presentation: {OUTPUT_PPTX}")

def console_log(msg):
    print(msg)
    sys.stdout.flush()

if __name__ == '__main__':
    create_presentation()
